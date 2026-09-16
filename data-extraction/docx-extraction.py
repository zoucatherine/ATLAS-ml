from docx import Document
from pathlib import Path

# Input file
input_file = Path("docx/BLS marketing plan 3.28.25 scrubbed.docx")

# Open DOCX
doc = Document(input_file)

markdown = []

# Extract paragraphs
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()

    if not text:
        continue

    # Preserve headings as Markdown headings
    if paragraph.style.name.startswith("Heading"):
        try:
            level = int(paragraph.style.name.split()[-1])
            markdown.append(f"{'#' * level} {text}")
        except ValueError:
            markdown.append(f"## {text}")
    else:
        markdown.append(text)

# Extract tables
for table in doc.tables:
    if not table.rows:
        continue

    rows = [
        [cell.text.strip() for cell in row.cells]
        for row in table.rows
    ]

    # First row becomes table header
    header = rows[0]

    markdown.append(
        "| " + " | ".join(header) + " |"
    )
    markdown.append(
        "| " + " | ".join(["---"] * len(header)) + " |"
    )

    # Remaining rows
    for row in rows[1:]:
        markdown.append(
            "| " + " | ".join(row) + " |"
        )

# Save Markdown file
output_file = Path("output.md")

with open(output_file, "w", encoding="utf-8") as f:
    f.write("\n\n".join(markdown))

print(f"Saved to {output_file}")